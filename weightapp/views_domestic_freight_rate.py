# -*- coding: utf-8 -*-
"""ตารางราคา ค่าขนส่งในประเทศ (issue #118) : หน้ารายการ / เพิ่ม / ลบ เส้นทาง + export Excel

ตาราง DomesticFreightRate / DomesticFreightRateTeam แยกจากส่งออก แต่ใช้หลักเดียวกัน
ยังไม่ผูกเมนู navbar เข้าด้วย URL ตรง ๆ (/domesticFreightRate/)
ยังไม่มีหน้าแก้ไข/ประวัติเวอร์ชัน และหน้ารายงานคิดเงินรายเดือน
"""
from datetime import date

import openpyxl
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Prefetch
from django.http import HttpResponse
from django.shortcuts import redirect, render
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from weightapp import views as core_views
from weightapp.models import (BaseCarTeam, BaseCompanyMapBaseCustomer, DomesticFreightRate,
                              DomesticFreightRateTeam, DomesticFreightWeightBand)
from weightapp.serializers import DomesticFreightRateSerializer

NO_NAME = '(ไม่พบชื่อ)'
BAND_ORDER = {DomesticFreightWeightBand.ALL: 0, DomesticFreightWeightBand.UNDER_31: 1,
              DomesticFreightWeightBand.FROM_31: 2}


def _sortedTeams(teams):
    """ทีมที่ระบุชื่อขึ้นก่อน (เรียงตามรหัส) แถว "ทุกทีม" ไว้ท้าย ในทีมเดียวกันเรียงตามช่วงน้ำหนัก"""
    return sorted(teams, key=lambda t: (t.team_id is None, t.team_id or '', BAND_ORDER.get(t.weight_band, 9)))


def _currentRates():
    """ใบที่ใช้อยู่ตอนนี้ เส้นทางละใบ พร้อมทีม"""
    return DomesticFreightRate.effectiveOn(
        date.today(),
        DomesticFreightRate.objects
        .select_related('origin', 'destination')
        .prefetch_related(Prefetch('teams', queryset=DomesticFreightRateTeam.objects.select_related('team'))))


@login_required(login_url='login')
def viewDomesticFreightRate(request):
    active = request.session.get('company_code')
    if not active:
        return redirect('logout')

    map_names = core_views._baseCompanyMapDisplayNameById()
    rows = _currentRates()
    for row in rows:
        row.sorted_teams = _sortedTeams(row.teams.all())
        row.has_specific_team = any(t.team_id for t in row.sorted_teams)
        row.team_group_count = len({t.team_id for t in row.sorted_teams})
        row.origin_display = map_names.get(row.origin_id) or NO_NAME
        row.destination_display = map_names.get(row.destination_id) or NO_NAME

    route_count = {}
    for row in rows:
        route_count[row.origin_id] = route_count.get(row.origin_id, 0) + 1
    origin_options = sorted(
        ({'id': oid, 'name': map_names.get(oid) or NO_NAME, 'count': n} for oid, n in route_count.items()),
        key=lambda o: core_views._thaiSortKey(o['name']))
    try:
        selected_origin = int(request.GET.get('origin') or '')
    except ValueError:
        selected_origin = None
    shown = rows if selected_origin is None else [r for r in rows if r.origin_id == selected_origin]

    groups = {}
    for row in shown:
        groups.setdefault(row.origin_id, []).append(row)
    groups = sorted(
        ({'origin_name': map_names.get(oid) or NO_NAME,
          'rows': sorted(g, key=lambda r: core_views._thaiSortKey(r.destination_display))}
         for oid, g in groups.items()),
        key=lambda g: core_views._thaiSortKey(g['origin_name']))

    context = {
        'dfr_page': 'active',
        'dfr_groups': groups,
        'dfr_origin_options': origin_options,
        'dfr_selected_origin': selected_origin,
        'dfr_route_total': len(rows),
        active: 'active',
    }
    return render(request, 'domesticFreightRate/viewDomesticFreightRate.html', context)


def _mapOptions():
    """ตัวเลือกต้นทาง (แถวที่ผูกบริษัท) และปลายทาง (แถวที่ติ๊ก "ใช้เป็นปลายทางในประเทศ")"""
    rows = core_views._baseCompanyMapRows()
    domestic_ids = set(BaseCompanyMapBaseCustomer.objects
                       .filter(is_domestic_destination=True).values_list('id', flat=True))
    origins = [{'id': r['id'], 'name': r['display']} for r in rows if r['base_company_id']]
    destinations = [{'id': r['id'], 'name': r['display']} for r in rows if r['id'] in domestic_ids]
    key = lambda o: core_views._thaiSortKey(o['name'] or '')
    return sorted(origins, key=key), sorted(destinations, key=key)


@login_required(login_url='login')
def createDomesticFreightRate(request):
    active = request.session.get('company_code')
    if not active:
        return redirect('logout')
    origins, destinations = _mapOptions()
    teams = [{'id': t['car_team_id'], 'name': t['car_team_name'] or t['car_team_id']}
             for t in BaseCarTeam.objects.order_by('car_team_id').values('car_team_id', 'car_team_name')]
    context = {
        'dfr_page': 'active',
        'dfr_origins': origins,
        'dfr_destinations': destinations,
        'dfr_teams': teams,
        'dfr_bands': [[value, label] for value, label in DomesticFreightWeightBand.choices],
        active: 'active',
    }
    return render(request, 'domesticFreightRate/createDomesticFreightRate.html', context)


# ระบุสิทธิ์เองทุกตัว : DEFAULT_PERMISSION_CLASSES ใน settings เขียนเป็น string ไม่ใช่ tuple ใช้ค่า default ไม่ได้
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def domesticFreightRateCreate(request):
    serializer = DomesticFreightRateSerializer(data=request.data, context={'request': request})
    if serializer.is_valid():
        rate = serializer.save()
        return Response({'id': rate.id}, status=status.HTTP_201_CREATED)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def domesticFreightRateDelete(request, id):
    """ลบเส้นทาง = ลบทุกเวอร์ชันของเส้นทางนั้น (ลำดับเดียวกับ internationalFreightRateDelete)"""
    obj = DomesticFreightRate.objects.filter(id=id).first()
    if obj is None:
        return Response({'error': 'ไม่พบข้อมูล'}, status=status.HTTP_404_NOT_FOUND)
    root_id = obj.root_id or obj.id
    with transaction.atomic():
        family = DomesticFreightRate.objects.filter(root_id=root_id)
        removed = family.count()
        family.exclude(id=root_id).delete()
        # ใบแรกชี้ตัวเอง ต้องตัดออกก่อน ไม่งั้น PROTECT จะกันการลบตัวเอง
        DomesticFreightRate.objects.filter(id=root_id).update(root=None)
        DomesticFreightRate.objects.filter(id=root_id).delete()
    return Response({'message': 'ลบสำเร็จ (ลบไป %s รายการ)' % removed}, status=status.HTTP_200_OK)


# ข้อความที่ผู้ใช้พิมพ์เอง (หมายเหตุ ชื่อทีม ชื่อในตาราง map) ห้ามลงไฟล์ดิบ ๆ
# openpyxl ถือว่าข้อความที่ขึ้นต้นด้วย = เป็นสูตร เปิดไฟล์แล้ว Excel จะรันสูตรนั้น (formula injection)
EXCEL_FORMULA_PREFIXES = ('=', '+', '-', '@', chr(9), chr(13))


def _excelText(value):
    if isinstance(value, str) and value.startswith(EXCEL_FORMULA_PREFIXES):
        return "'" + value
    return value


def _teamColumnLabel(team_row):
    team = (team_row.team.car_team_name or team_row.team_id) if team_row.team_id else 'ทุกทีม'
    if team_row.weight_band == DomesticFreightWeightBand.ALL:
        return team
    return '%s\n%s' % (team, team_row.get_weight_band_display())


@login_required(login_url='login')
def exportExcelDomesticFreightRate(request):
    """ตารางราคาที่ใช้อยู่ปัจจุบัน หน้าตาคล้ายชีต "อัตราค่าขนส่ง ใช้งาน" ของไฟล์เดิม

    คอลัมน์ราคาแตกตาม (ทีม, ช่วงน้ำหนัก) ที่มีจริง เช่น ROI | ทุกทีม 0-30.99 | ทุกทีม 31 ขึ้นไป
    ชีตที่ 2 เป็นรายการทีมแบบแถวเรียบ มี % สูญเสียและหมายเหตุของแต่ละแถว
    """
    if not request.session.get('company_code'):
        return redirect('logout')
    map_names = core_views._baseCompanyMapDisplayNameById()
    rates = _currentRates()
    rates.sort(key=lambda r: (core_views._thaiSortKey(map_names.get(r.origin_id) or ''),
                              r.distance if r.distance is not None else 0,
                              core_views._thaiSortKey(map_names.get(r.destination_id) or '')))

    # คอลัมน์ราคา : ทีมที่ระบุชื่อก่อน แล้วค่อย "ทุกทีม" ในทีมเดียวกันเรียงตามช่วงน้ำหนัก
    columns = {}
    for rate in rates:
        for t in rate.teams.all():
            key = (t.team_id is None, t.team_id or '', BAND_ORDER.get(t.weight_band, 9))
            columns.setdefault(key, _teamColumnLabel(t))
    column_keys = sorted(columns)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'อัตราค่าขนส่ง'
    ws.append(['อัตราค่าขนส่ง ในประเทศ (ไม่รวมภาษีมูลค่าเพิ่ม) ณ วันที่ %s' % date.today().strftime('%d/%m/%Y')])
    ws.append([])
    header = ['ที่', 'ต้นทาง', 'ลูกค้า (ปลายทาง)', 'ระยะทาง (กม.)'] + [_excelText(columns[k]) for k in column_keys] + [
        'ราคาน้ำมันฐาน (บาท/ลิตร)', 'ขั้นปรับน้ำมัน', 'วันที่เริ่มใช้', 'หมายเหตุ']
    ws.append(header)
    header_row = ws.max_row
    for no, rate in enumerate(rates, start=1):
        price = {(t.team_id is None, t.team_id or '', BAND_ORDER.get(t.weight_band, 9)): t.freight_rate
                 for t in rate.teams.all()}
        ws.append([no, _excelText(map_names.get(rate.origin_id)), _excelText(map_names.get(rate.destination_id)),
                   float(rate.distance) if rate.distance is not None else None]
                  + [float(price[k]) if price.get(k) is not None else None for k in column_keys]
                  + [rate.baseFuelPriceLabel(), rate.fuelStepLabel(),
                     rate.effective_date, _excelText(rate.note)])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ws['A1'].font = Font(bold=True, size=13)
    widths = [6, 22, 34, 12] + [14] * len(column_keys) + [18, 30, 14, 30]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=header_row + 1, min_col=len(widths) - 1, max_col=len(widths) - 1):
        for cell in row:
            cell.number_format = 'DD/MM/YYYY'

    detail = wb.create_sheet('รายละเอียดทีม')
    detail.append(['ต้นทาง', 'ลูกค้า (ปลายทาง)', 'ทีม', 'ช่วงน้ำหนัก', 'อัตราค่าขนส่ง (บาท/ตัน)',
                   'หักค่าสูญเสีย (%)', 'หมายเหตุ'])
    for rate in rates:
        for t in _sortedTeams(rate.teams.all()):
            detail.append([_excelText(map_names.get(rate.origin_id)), _excelText(map_names.get(rate.destination_id)),
                           _excelText((t.team.car_team_name or t.team_id) if t.team_id else 'ทุกทีม'),
                           t.get_weight_band_display(), float(t.freight_rate), float(t.loss_percent),
                           _excelText(t.note)])
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate([22, 34, 22, 22, 18, 16, 30], start=1):
        detail.column_dimensions[get_column_letter(i)].width = width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="domestic_freight_rate_%s.xlsx"' % (
        date.today().strftime('%Y%m%d'))
    wb.save(response)
    return response
