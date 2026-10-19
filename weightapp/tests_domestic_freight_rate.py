# -*- coding: utf-8 -*-
"""ตารางราคาค่าขนส่งในประเทศ (issue #118) : DomesticFreightRate / DomesticFreightRateTeam

หน้ารายการ / เพิ่ม / ลบ เส้นทาง + export Excel
"""
import datetime
import io
from decimal import Decimal as D

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse

from weightapp.models import (INTERNATIONAL_FREIGHT_RATE_FIRST_DATE, BaseCarTeam, BaseCompany,
                              BaseCompanyMapBaseCustomer, BaseCustomer, DomesticFreightRate,
                              DomesticFreightRateTeam, InternationalFreightRate, UserProfile)


# ปิด whitenoise manifest ใน test เพราะไม่ได้รัน collectstatic
@override_settings(STATICFILES_STORAGE='django.contrib.staticfiles.storage.StaticFilesStorage')
class DomesticFreightRateTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.company = BaseCompany.objects.create(name='ศิลาชัย', code='SLC')
        BaseCustomer.objects.create(customer_id='CP-1', customer_name='ปูนซิเมนต์ไทย/พุนพิน')
        BaseCustomer.objects.create(customer_id='PORT-1', customer_name='ท่าเรือส่งออก')
        cls.origin = BaseCompanyMapBaseCustomer.objects.create(name='ศิลาชัย', base_company=cls.company)
        cls.dest = BaseCompanyMapBaseCustomer.objects.create(
            name='CPAC พุนพิน', base_customer_id='CP-1', is_domestic_destination=True)
        cls.export_dest = BaseCompanyMapBaseCustomer.objects.create(
            name='ท่าเรือส่งออก', base_customer_id='PORT-1', is_export_destination=True)
        cls.roi = BaseCarTeam.objects.create(car_team_id='ROI', car_team_name='ร้อยเกาะ')
        cls.user = User.objects.create_user('u', 'u@t.com', 'pw12345!')
        UserProfile.objects.create(user=cls.user).company.add(cls.company)

    def setUp(self):
        self.client.force_login(self.user)
        session = self.client.session
        session['company_code'] = 'SLC'
        session.save()

    def payload(self, **over):
        data = {
            'origin': self.origin.pk, 'destination': self.dest.pk, 'distance': '36',
            'base_fuel_price': '32.00', 'base_fuel_price_max': '32.99',
            'fuel_step_price': '0.50', 'fuel_step_amount': '0.52',
            'note': 'น้ำมัน + 0.50 บาท ปรับ + 0.52',
            'teams': [
                {'team_id': 'ROI', 'weight_band': 'all', 'freight_rate': '81.48', 'loss_percent': '0'},
                {'team_id': None, 'weight_band': 'under_31', 'freight_rate': '76.48', 'loss_percent': '1.5'},
                {'team_id': None, 'weight_band': 'from_31', 'freight_rate': '71.48', 'loss_percent': '1.5'},
            ],
        }
        data.update(over)
        return data

    def create(self, **over):
        return self.client.post(reverse('domesticFreightRateCreate'), self.payload(**over),
                                content_type='application/json')

    # ---------- หน้าเว็บ ----------

    def test_pages_render(self):
        self.assertEqual(self.client.get(reverse('viewDomesticFreightRate')).status_code, 200)
        res = self.client.get(reverse('createDomesticFreightRate'))
        # ชื่อที่แสดงใช้กติกาเดียวกับส่งออก : ชื่อลูกค้า -> ชื่อบริษัท -> ชื่อในตาราง map
        self.assertContains(res, 'ปูนซิเมนต์ไทย/พุนพิน')
        # ปลายทางเลือกได้เฉพาะแถวที่ติ๊กในประเทศ
        self.assertNotContains(res, 'ท่าเรือส่งออก')

    def test_list_shows_route_and_teams(self):
        self.assertEqual(self.create().status_code, 201)
        res = self.client.get(reverse('viewDomesticFreightRate'))
        self.assertContains(res, 'ปูนซิเมนต์ไทย/พุนพิน')
        self.assertContains(res, '81.48')
        self.assertContains(res, 'ทุก 0.50 บาท/ลิตร ปรับ 0.52 บาท/ตัน')

    # ---------- สร้าง ----------

    def test_create_saves_header_and_teams(self):
        res = self.create()
        self.assertEqual(res.status_code, 201, res.content)
        rate = DomesticFreightRate.objects.get()
        self.assertEqual((rate.version, rate.root_id), (1, rate.id))
        self.assertEqual(rate.effective_date, INTERNATIONAL_FREIGHT_RATE_FIRST_DATE)
        self.assertEqual(rate.user_created, self.user)
        self.assertEqual(rate.teams.count(), 3)
        team = rate.teams.get(team=None, weight_band='under_31')
        self.assertEqual((team.freight_rate, team.loss_percent), (D('76.48'), D('1.5')))
        # ตารางของส่งออกไม่ถูกแตะ
        self.assertFalse(InternationalFreightRate.objects.exists())

    def test_create_requires_origin_destination_and_team(self):
        res = self.create(origin=None, destination=None, teams=[])
        self.assertEqual(res.status_code, 400)
        for key in ('origin', 'destination', 'teams'):
            self.assertIn(key, res.json())

    def test_destination_must_be_flagged_domestic(self):
        res = self.create(destination=self.export_dest.pk)
        self.assertEqual(res.status_code, 400)
        self.assertIn('destination', res.json())

    def test_duplicate_route_rejected(self):
        self.create()
        res = self.create()
        self.assertEqual(res.status_code, 400)
        self.assertEqual(DomesticFreightRate.objects.count(), 1)

    def test_all_band_cannot_mix_with_split_bands(self):
        res = self.create(teams=[
            {'team_id': None, 'weight_band': 'all', 'freight_rate': '70'},
            {'team_id': None, 'weight_band': 'from_31', 'freight_rate': '65'},
        ])
        self.assertEqual(res.status_code, 400)

    def test_same_band_twice_rejected(self):
        res = self.create(teams=[
            {'team_id': 'ROI', 'weight_band': 'under_31', 'freight_rate': '70'},
            {'team_id': 'ROI', 'weight_band': 'under_31', 'freight_rate': '65'},
        ])
        self.assertEqual(res.status_code, 400)

    def test_fuel_step_needs_both_values(self):
        self.assertEqual(self.create(fuel_step_amount=None).status_code, 400)
        res = self.create(fuel_step_price=None, fuel_step_amount=None)
        self.assertEqual(res.status_code, 201, res.content)

    def test_fuel_base_upper_below_lower_rejected(self):
        self.assertEqual(self.create(base_fuel_price_max='31.00').status_code, 400)

    def test_loss_percent_out_of_range_rejected(self):
        res = self.create(teams=[{'team_id': None, 'weight_band': 'all', 'freight_rate': '70',
                                  'loss_percent': '120'}])
        self.assertEqual(res.status_code, 400)

    def test_effective_date_kept_when_given(self):
        self.create(effective_date='2026-08-22')
        self.assertEqual(DomesticFreightRate.objects.get().effective_date, datetime.date(2026, 8, 22))

    def test_api_requires_login(self):
        self.client.logout()
        self.assertIn(self.create().status_code, (401, 403))

    # ---------- ลบ ----------

    def test_delete_removes_whole_family(self):
        self.create()
        root = DomesticFreightRate.objects.get()
        v2 = DomesticFreightRate.objects.create(origin=self.origin, destination=self.dest,
                                                root=root, version=2, effective_date=datetime.date(2026, 9, 1))
        DomesticFreightRateTeam.objects.create(rate=v2, freight_rate=D('90'))
        res = self.client.delete(reverse('domesticFreightRateDelete', args=[v2.pk]))
        self.assertEqual(res.status_code, 200)
        self.assertFalse(DomesticFreightRate.objects.exists())
        self.assertFalse(DomesticFreightRateTeam.objects.exists())

    def test_delete_missing_returns_404(self):
        self.assertEqual(self.client.delete(reverse('domesticFreightRateDelete', args=[999])).status_code, 404)

    # ---------- Excel ----------

    def test_excel_export_has_team_band_columns(self):
        self.create()
        res = self.client.get(reverse('exportExcelDomesticFreightRate'))
        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb['อัตราค่าขนส่ง']
        header = [c.value for c in ws[3]]
        self.assertEqual(header[:4], ['ที่', 'ต้นทาง', 'ลูกค้า (ปลายทาง)', 'ระยะทาง (กม.)'])
        # ทีมที่ระบุชื่อก่อน แล้วค่อยทุกทีมแยกช่วง
        self.assertEqual(header[4:7], ['ร้อยเกาะ', 'ทุกทีม\nน้ำหนัก 0 - 30.99 ตัน', 'ทุกทีม\nน้ำหนัก 31 ตันขึ้นไป'])
        row = [c.value for c in ws[4]]
        self.assertEqual(row[1:7], ['ศิลาชัย', 'ปูนซิเมนต์ไทย/พุนพิน', 36.0, 81.48, 76.48, 71.48])
        self.assertEqual(wb['รายละเอียดทีม'].max_row, 4)      # หัว + 3 แถวทีม
